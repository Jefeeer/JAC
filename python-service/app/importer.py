"""Supplier price-list import (CSV or Excel) → parts upsert.

Mirrors the CSV fallback in server/actions/admin.ts so both paths accept the
same template. Parsing is pure (testable without a database); `upsert_parts`
does the I/O.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable

MAX_ROWS = 2000

# Accept common supplier header spellings.
HEADER_ALIASES: dict[str, str] = {
    "part_no": "part_number",
    "part_#": "part_number",
    "partnumber": "part_number",
    "sku": "part_number",
    "description": "name",
    "part_name": "name",
    "qty": "stock_qty",
    "quantity": "stock_qty",
    "stock": "stock_qty",
    "on_hand": "stock_qty",
    "srp": "price",
    "unit_price": "price",
    "oem": "oem_number",
    "oem_no": "oem_number",
    "models": "compatible_models",
    "fits": "compatible_models",
    "reorder": "reorder_level",
    "lead_time": "lead_time_days",
}


class ImportError_(ValueError):
    """Raised for file-level problems (bad type, missing columns, too many rows)."""


@dataclass
class PartRow:
    part_number: str
    name: str
    category_slug: str
    price: Decimal | None
    stock_qty: int
    reorder_level: int
    oem_number: str | None = None
    brand: str = "JAC Genuine"
    unit: str = "pc"
    summary: str | None = None
    lead_time_days: int | None = None
    compatible_models: list[str] = field(default_factory=list)


@dataclass
class ParseResult:
    rows: list[PartRow]
    skipped: list[str]


def normalize_header(h: Any) -> str:
    key = re.sub(r"[^a-z0-9#]+", "_", str(h or "").strip().lower()).strip("_")
    return HEADER_ALIASES.get(key, key)


def slugify(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def _money(v: Any) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(Decimal("0.01"))
    s = re.sub(r"[,\s₱]|PHP", "", str(v), flags=re.I)
    if not s:
        return None
    try:
        d = Decimal(s)
    except InvalidOperation as e:
        raise ValueError("price") from e
    if d < 0:
        raise ValueError("price")
    return d.quantize(Decimal("0.01"))


def _int(v: Any, default: int) -> int:
    if v is None or str(v).strip() == "":
        return default
    n = float(str(v).replace(",", ""))
    if n != n or n in (float("inf"), float("-inf")):  # NaN / inf
        raise ValueError("number")
    return max(0, round(n))


def _text(v: Any) -> str | None:
    s = str(v).strip() if v is not None else ""
    return s or None


def read_table(filename: str, data: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    """Read a CSV / XLSX upload into normalized headers and row dicts."""
    name = filename.lower()
    if name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text))
        table = [r for r in reader]
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        raise ImportError_("Upload a .csv or .xlsx file.")

    table = [r for r in table if any(c not in (None, "") for c in r)]
    if not table:
        raise ImportError_("The file is empty.")
    headers = [normalize_header(h) for h in table[0]]
    rows = [{h: (r[i] if i < len(r) else None) for i, h in enumerate(headers) if h} for r in table[1:]]
    return headers, rows


def parse_parts(filename: str, data: bytes) -> ParseResult:
    headers, raw = read_table(filename, data)
    if "part_number" not in headers or "name" not in headers:
        raise ImportError_("The file needs at least part_number and name columns.")
    if len(raw) > MAX_ROWS:
        raise ImportError_(f"Import up to {MAX_ROWS:,} rows at a time.")

    rows: list[PartRow] = []
    skipped: list[str] = []
    seen: set[str] = set()
    for i, r in enumerate(raw, start=2):  # row 1 is the header
        pn = (_text(r.get("part_number")) or "").upper()
        name = _text(r.get("name"))
        if not pn or not name:
            skipped.append(f"Row {i}: {pn or '(no part number)'}")
            continue
        if pn in seen:
            skipped.append(f"Row {i}: {pn} (duplicate in file)")
            continue
        try:
            row = PartRow(
                part_number=pn,
                name=name[:200],
                category_slug=slugify(_text(r.get("category")) or "engine-filtration"),
                price=_money(r.get("price")),
                stock_qty=_int(r.get("stock_qty"), 0),
                reorder_level=_int(r.get("reorder_level"), 5),
                oem_number=_text(r.get("oem_number")),
                brand=_text(r.get("brand")) or "JAC Genuine",
                unit=_text(r.get("unit")) or "pc",
                summary=_text(r.get("summary")),
                lead_time_days=_int(r.get("lead_time_days"), -1) if _text(r.get("lead_time_days")) else None,
                compatible_models=[m.strip() for m in re.split(r"[;|]", _text(r.get("compatible_models")) or "") if m.strip()],
            )
        except ValueError as e:
            skipped.append(f"Row {i}: {pn} (invalid {e})")
            continue
        seen.add(pn)
        rows.append(row)
    return ParseResult(rows=rows, skipped=skipped)


def _record(row: PartRow, category_id: str | None) -> dict[str, Any]:
    return {
        "part_number": row.part_number,
        "name": row.name,
        "category_id": category_id,
        "price": str(row.price) if row.price is not None else None,
        "price_on_request": row.price is None,
        "stock_qty": row.stock_qty,
        "reorder_level": row.reorder_level,
        "oem_number": row.oem_number,
        "brand": row.brand,
        "unit": row.unit,
        "summary": row.summary,
        "lead_time_days": row.lead_time_days,
    }


def upsert_parts(db: Any, rows: Iterable[PartRow]) -> dict[str, int]:
    """Insert new part numbers, update existing ones. Fitment is replaced
    only when the file provides compatible_models for that row."""
    rows = list(rows)
    cats = db.table("part_categories").select("id, slug").execute().data or []
    cat_ids = {c["slug"]: c["id"] for c in cats}

    existing: dict[str, dict[str, Any]] = {}
    numbers = [r.part_number for r in rows]
    for i in range(0, len(numbers), 200):
        res = db.table("parts").select("id, slug, part_number").in_("part_number", numbers[i : i + 200]).execute()
        existing.update({p["part_number"]: p for p in res.data or []})

    inserted = updated = 0
    for row in rows:
        rec = _record(row, cat_ids.get(row.category_slug))
        if row.part_number in existing:
            part_id = existing[row.part_number]["id"]
            db.table("parts").update(rec).eq("id", part_id).execute()
            updated += 1
        else:
            rec["slug"] = slugify(f"{row.name}-{row.part_number}")[:120]
            rec["is_published"] = True
            part_id = db.table("parts").insert(rec).execute().data[0]["id"]
            inserted += 1
        if row.compatible_models:
            db.table("part_compatibility").delete().eq("part_id", part_id).execute()
            db.table("part_compatibility").insert([{"part_id": part_id, "model": m} for m in dict.fromkeys(row.compatible_models)]).execute()
    return {"inserted": inserted, "updated": updated}
